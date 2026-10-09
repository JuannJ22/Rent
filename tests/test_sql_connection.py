from contextlib import contextmanager
from datetime import date
from types import SimpleNamespace
import sys

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from hojas import hoja01_loader as loader
from rentabilidad.infra.sql_server import SqlServerConfig, dataframe_reader


def test_siigo_sql_text_repairs_both_encodings_without_changing_amounts():
    original = pd.DataFrame({
        "Descripcion": ["CU" + chr(129) + "ETE", "CU" + chr(169) + "ETE", "CUÑETE", "CEMENTO", None],
        "Ventas": [10.5, -10.5, 0, 15, 2],
    })
    result = loader._normalize_sql_text(original)
    assert result["Descripcion"].tolist() == ["CUÑETE", "CUÑETE", "CUÑETE", "CEMENTO", None]
    assert result["Ventas"].equals(original["Ventas"])
    assert chr(129) in original.loc[0, "Descripcion"]


def test_odbc_quotes_password_and_explicit_tls_flags():
    config = SqlServerConfig("server", "db", user="reader", password="a;PWD=b}c",
                             encrypt=True, trust_server_certificate=False)
    value = config.connection_string()
    assert "PWD={a;PWD=b}}c}" in value
    assert "Encrypt=yes" in value
    assert "TrustServerCertificate=no" in value


def test_reader_closes_cursor_and_connection_on_query_failure(monkeypatch):
    events = []
    class Cursor:
        def execute(self, *args):
            raise RuntimeError("query failed")
        def close(self):
            events.append("cursor closed")
    class Connection:
        def cursor(self):
            return Cursor()
        def close(self):
            events.append("connection closed")
    monkeypatch.setitem(sys.modules, "pyodbc", SimpleNamespace(connect=lambda *a, **k: Connection()))
    with pytest.raises(RuntimeError, match="query failed"):
        with dataframe_reader(SqlServerConfig("server", "db")) as read:
            read("SELECT 1")
    assert events == ["cursor closed", "connection closed"]


def test_all_report_queries_share_connection_and_configured_databases(monkeypatch):
    calls = []
    connections = []
    @contextmanager
    def reader(config):
        connections.append(config)
        def read(query, params=None):
            calls.append((query, params))
            return pd.DataFrame()
        yield read
    monkeypatch.setattr(loader, "dataframe_reader", reader)
    loader._load_sql_report_data(SqlServerConfig("server", "db"), date(2026, 10, 5), {
        "SQL_RENT_DATABASE": "Rent]DB", "SQL_CAT_DATABASE": "Catalog",
        "SQL_MOV_DATABASE": "Movements", "SQL_REQUIRE_PRODUCT_SNAPSHOT": False,
    })
    assert len(connections) == 1
    assert len(calls) == 19
    dated = [(q, p) for q, p in calls if "WHERE FECHA" in q]
    assert len(dated) == 14
    assert all(p == ["2026-10-05"] and "[Rent]]DB]" in q for q, p in dated)
    assert any("[Catalog]" in q for q, _ in calls)
    movement_queries = [(q, p) for q, p in calls if "vw_movimientos_vendedores_informe" in q]
    assert len(movement_queries) == 1
    assert movement_queries[0][1] == ["2026-10-05"]
    assert "WHERE FechaDctoMov" in movement_queries[0][0]


def test_no_sql_rows_fails_before_touching_excel(monkeypatch):
    monkeypatch.setenv("SQL_SERVER", "server")
    monkeypatch.setenv("SQL_DATABASE", "db")
    monkeypatch.setenv("SQL_USER", "reader")
    empty = pd.DataFrame(columns=["NIT - SUCURSAL - CLIENTE", "VENTAS", "COSTO"])
    monkeypatch.setattr(loader, "_load_sql_report_data", lambda *args: (
        empty, empty, {}, {}, empty, empty, empty,
    ))
    monkeypatch.setattr(sys, "argv", ["loader", "--sql", "--require-sql-data",
                                     "--excel", "must-not-open.xlsx", "--fecha", "2026-10-05"])
    with pytest.raises(SystemExit) as error:
        loader.main()
    assert error.value.code == 36


def test_sql_loader_writes_real_excel_without_excz_files(tmp_path, monkeypatch):
    path = tmp_path / "Octubre 05 SQL corregido-20261007-154446.xlsx"
    book = Workbook()
    headers = ["NIT", "NIT - SUCURSAL - CLIENTE", "CANTIDAD", "COD. VENDEDOR",
               "VENTAS", "COSTO", "% RENTA.", "% UTILI."]
    for column, name in enumerate(headers, start=1):
        book.active.cell(6, column, name)
    book.save(path)
    data = pd.DataFrame([{
        "NIT - SUCURSAL - CLIENTE": "900123 - 1 - CLIENTE",
        "CANTIDAD": 2, "VENTAS": 1000, "COSTO": 600,
        "COD. VENDEDOR": 30, "DOCUMENTO": "F202 91658",
        "% RENTA.": 0.4, "% UTILI.": 2 / 3,
    }])
    empty = pd.DataFrame()
    monkeypatch.setenv("SQL_SERVER", "server")
    monkeypatch.setenv("SQL_DATABASE", "db")
    monkeypatch.setenv("SQL_USER", "reader")
    third_parties = pd.DataFrame([{"NitNit": "900123", "VendedorNit": 24, "PrecioNit": 1}])
    monkeypatch.setattr(loader, "_load_sql_report_data", lambda *args: (
        data, empty, {}, {}, empty, third_parties, empty,
    ))
    monkeypatch.setattr(sys, "argv", [
        "loader", "--sql", "--require-sql-data", "--excel", str(path),
        "--fecha", "2026-10-05", "--exczdir", str(tmp_path / "absent"),
    ])
    loader.main()
    result = load_workbook(path)
    sheet = result.active
    assert sheet["E7"].value == 1000
    assert sheet["F7"].value == 600
    assert sheet["C7"].value == 2
    assert str(sheet["D7"].value) == "30"
    assert sheet.title == "OCTUBRE 05"
    result.close()


def test_json_false_flags_override_true_environment(monkeypatch):
    monkeypatch.setenv("SQL_TRUSTED", "1")
    monkeypatch.setenv("SQL_ENCRYPT", "1")
    monkeypatch.setattr(sys, "argv", ["loader"])
    args = SimpleNamespace(sql_server="server", sql_database="db", sql_user="reader",
                           sql_password=None, sql_driver=None, sql_trusted=False,
                           sql_config_data={"SQL_TRUSTED": False, "SQL_ENCRYPT": False})
    config = loader._build_sql_config(args)
    assert config.trusted_connection is False
    assert config.encrypt is False


def test_gui_accepts_sql_environment_without_creating_config(tmp_path, monkeypatch):
    from rentabilidad.app.dto import GenerarInformeRequest
    from rentabilidad.app.use_cases import generar_informe_automatico as automatic
    from rentabilidad.infra.logging_bus import EventBus

    monkeypatch.setenv("SQL_SERVER", "server")
    monkeypatch.setenv("SQL_DATABASE", "db")
    monkeypatch.setattr(automatic.settings, "sql_config", None)
    monkeypatch.setattr(automatic, "_ensure_sql_config_template", lambda: pytest.fail("Unexpected config creation"))
    result = automatic.run(GenerarInformeRequest(
        ruta_plantilla=str(tmp_path / "missing.xlsx"), usar_sql=True,
    ), EventBus())
    assert result.ok is False
    assert "No existe la plantilla" in result.mensaje


def test_daily_excel_preserves_sql_total_and_compacts_customers_repeatably():
    book = Workbook()
    main = book.active
    main['A7'] = 900123
    zone = book.create_sheet('CCOSTO 4')
    zone.append(['ZONA', 'DESCRIPCION', 'CANTIDAD', 'VENTAS', 'COSTO', '% RENTA', '% UTIL'])
    zone.append(['4', 'PRODUCTO', 10, 100, 60, 40, 66.67])
    zone.append(['4', 'DEVOLUCION', -2, -20, -12, 40, 66.67])
    zone.append(['TOTAL TIENDA PINTUCO (ZONA 7)', '', 8, 80, 48, 40, 66.67])
    third = book.create_sheet('TERCEROS')
    third.sheet_state = 'hidden'
    third.append([900123, 1, 24])
    third.append([900999, 12, 30])
    third.append([900123, 6, 26])
    loader._finalize_daily_sql_workbook(book, main)
    assert list(zone.values)[-1] == ('TOTAL TIENDA PINTUCO (ZONA 7)', '', 8, 80, 48, 40, 66.67)
    assert list(third.values) == [(900123, 6, 26)]
    assert third.sheet_state == 'hidden'
    loader._finalize_daily_sql_workbook(book, main)
    assert zone.max_row == 4
    assert third.max_row == 1
