from datetime import date
from pathlib import Path
from types import SimpleNamespace
import pandas as pd
import pytest
from openpyxl import Workbook
from servicios import etl_rentabilidad as runner


def add_snapshot(base, day):
    from rentabilidad.infra import product_snapshots as snapshots
    from rentabilidad.infra.sql_server import SqlServerConfig
    from unittest.mock import patch
    from datetime import datetime
    with patch.object(snapshots, 'fetch_dataframe', return_value=pd.DataFrame({'DESCRIPCION': ['PRODUCTO'], 'Precio1': [100]})):
        snapshots.capture_snapshot(SqlServerConfig('server', 'SiigoRent'), base / 'Productos',
            now=datetime(day.year, day.month, day.day, 21, tzinfo=snapshots.COLOMBIA))


def write_report(path, day):
    wb = Workbook()
    wb.active.title = f'OCTUBRE {day.day:02d}'
    wb.active.cell(7, 6, 100)
    wb.active.cell(7, 7, 60)
    wb.save(path)


def test_etl_publishes_validated_report_and_reuses_it(tmp_path, monkeypatch):
    template = tmp_path / 'template.xlsx'
    Workbook().save(template)
    day = date(2026, 10, 6)
    add_snapshot(tmp_path, day)
    calls = []
    monkeypatch.setattr(runner, 'check_sql', lambda day: calls.append(day))
    def run(command, **kwargs):
        assert '--fecha' in command and '2026-10-06' in command
        assert '--require-sql-data' not in command  # Compatible con la instalacion anterior.
        write_report(Path(command[command.index('--excel') + 1]), day)
        return SimpleNamespace(returncode=0, stdout='OK', stderr='')
    monkeypatch.setattr(runner.subprocess, 'run', run)
    out = runner.generate(tmp_path, template, day, 60)
    before = out.read_bytes()
    assert out.name == 'Octubre 06.xlsx'
    assert runner.generate(tmp_path, template, day, 60) == out
    assert out.read_bytes() == before and calls == [day]


@pytest.mark.parametrize('failure', ['exit', 'invalid'])
def test_etl_failure_does_not_replace_existing_report(tmp_path, monkeypatch, failure):
    template = tmp_path / 'template.xlsx'
    Workbook().save(template)
    day = date(2026, 10, 6)
    out = tmp_path / 'InformesDiarios/2026/Octubre/Octubre 06.xlsx'
    out.parent.mkdir(parents=True)
    write_report(out, day)
    before = out.read_bytes()
    add_snapshot(tmp_path, day)
    monkeypatch.setattr(runner, 'check_sql', lambda _: None)
    monkeypatch.setattr(runner.subprocess, 'run', lambda *a, **k: SimpleNamespace(
        returncode=1 if failure == 'exit' else 0, stdout='', stderr='fallo'))
    with pytest.raises(RuntimeError):
        runner.generate(tmp_path, template, day, 60, replace=True)
    assert out.read_bytes() == before


def test_etl_requires_zone_total_when_detail_exists(tmp_path):
    path = tmp_path / 'test.xlsx'
    day = date(2026, 10, 6)
    write_report(path, day)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    zone = wb.create_sheet('CCOSTO 4')
    zone.append(['CENTRO', 'DESCRIPCION', 'CANTIDAD', 'VENTAS', 'COSTOS'])
    zone.append(['4', 'PRODUCTO', 1, 100, 60])
    wb.save(path)
    with pytest.raises(RuntimeError, match='Falta el total'):
        runner.validate_report(path, day)


def test_historical_report_never_uses_current_prices(tmp_path, monkeypatch):
    template = tmp_path / 'template.xlsx'
    Workbook().save(template)
    monkeypatch.setattr(runner, 'check_sql', lambda _: pytest.fail('No debe leer SQL sin la captura requerida'))
    with pytest.raises(RuntimeError, match='Falta la copia de productos'):
        runner.generate(tmp_path, template, date(2026, 10, 6), 60)
