from zipfile import ZipFile, ZIP_DEFLATED
from rentabilidad.infra.excel_compat import normalize_excel_application


def test_application_change_preserves_all_other_parts_and_is_idempotent(tmp_path):
    path = tmp_path / 'informe.xlsx'
    original = {
        'docProps/app.xml': b'<Properties><Application>Microsoft Excel Compatible / Openpyxl 3.1.5</Application><AppVersion>3.1</AppVersion></Properties>',
        'xl/worksheets/sheet1.xml': '<worksheet>Ñ =SUM(A1:A3)</worksheet>'.encode(),
        'xl/styles.xml': b'<styleSheet/>',
    }
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, data in original.items(): z.writestr(name, data)
    assert normalize_excel_application(path)
    with ZipFile(path) as z:
        for name, data in original.items():
            assert z.read(name) == (data.replace(b'Microsoft Excel Compatible / Openpyxl 3.1.5', b'Microsoft Excel') if name == 'docProps/app.xml' else data)
    assert not normalize_excel_application(path)


def test_report_presentation_numeric_vendor_left_alignment_and_hidden_origin():
    from openpyxl import Workbook
    from rentabilidad.infra.excel_compat import format_report_presentation
    wb = Workbook()
    ws = wb.active
    ws.append(['COD. VENDEDOR', 'CANTIDAD'])
    ws.append(['25', 3])
    ws.append(['26.0', 4])
    ws.append(['=VLOOKUP(A1,B1:C2,2,0)', 5])
    wb.create_sheet('ORIGEN_PRECIOS')
    format_report_presentation(wb)
    assert ws['A2'].value == 25 and ws['A2'].data_type == 'n'
    assert ws['A3'].value == 26
    assert ws['A2'].alignment.horizontal == 'left'
    assert ws['A4'].data_type == 'f'
    assert ws['B2'].value == 3
    assert wb['ORIGEN_PRECIOS'].sheet_state == 'hidden'
