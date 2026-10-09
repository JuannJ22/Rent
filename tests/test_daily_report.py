from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from rentabilidad import daily_report as daily


def test_yesterday_uses_colombia_even_when_utc_has_changed_day(monkeypatch):
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 7, 2, tzinfo=timezone.utc).astimezone(tz)
    monkeypatch.setattr(daily, "datetime", Clock)
    assert daily.target_date(None, "America/Bogota") == date(2026, 10, 5)


def test_failed_loader_preserves_existing_report_and_cleans_stage(tmp_path, monkeypatch):
    template = tmp_path / "PLANTILLA.xlsx"
    Workbook().save(template)
    destination = tmp_path / "InformesDiarios/2026/Octubre/Octubre 05.xlsx"
    destination.parent.mkdir(parents=True)
    destination.write_bytes(b"previous report")

    def fail(command, **kwargs):
        staged = Path(command[command.index("--excel") + 1])
        staged.write_bytes(b"incomplete")
        return SimpleNamespace(returncode=36, stdout="SQL sin datos", stderr="")
    monkeypatch.setattr(daily.subprocess, "run", fail)
    with pytest.raises(RuntimeError):
        daily.generate_report(template, destination, date(2026, 10, 5),
                              sql_config=None, timeout=10, replace=True)
    assert destination.read_bytes() == b"previous report"
    assert not list(destination.parent.glob(".daily-*"))


def test_success_publishes_valid_workbook_and_skips_repeat(tmp_path, monkeypatch):
    template = tmp_path / "PLANTILLA.xlsx"
    Workbook().save(template)
    destination = tmp_path / "InformesDiarios/2026/Octubre/Octubre 05.xlsx"
    calls = []

    def success(command, **kwargs):
        calls.append(command)
        assert "--sql" in command and "--require-sql-data" in command
        assert kwargs["timeout"] == 10
        staged = Path(command[command.index("--excel") + 1])
        assert staged.name == destination.name
        book = load_workbook(staged)
        book.active["A1"] = "SQL report"
        book.save(staged)
        return SimpleNamespace(returncode=0, stdout="OK", stderr="")
    monkeypatch.setattr(daily.subprocess, "run", success)
    for _ in range(2):
        daily.generate_report(template, destination, date(2026, 10, 5),
                              sql_config=None, timeout=10, replace=False)
    assert len(calls) == 1
    book = load_workbook(destination)
    assert book.active["A1"].value == "SQL report"
    book.close()


def test_concurrent_execution_rejected_and_lock_released(tmp_path):
    path = tmp_path / "report.lock"
    with daily.report_lock(path):
        with pytest.raises(RuntimeError, match="otra ejecución"):
            with daily.report_lock(path):
                pytest.fail("Second execution acquired the lock")
    with daily.report_lock(path):
        pass


def test_check_sql_requires_no_template_and_never_creates_report(tmp_path, monkeypatch):
    calls = []
    def run(command, **kwargs):
        calls.append(command)
        return SimpleNamespace(returncode=0, stdout="OK", stderr="")
    monkeypatch.setattr(daily.subprocess, "run", run)
    daily.generate_report(tmp_path / "missing.xlsx", tmp_path / "out.xlsx",
                          date(2026, 10, 5), sql_config="sql.json", timeout=10,
                          replace=False, check_only=True)
    assert "--check-sql" in calls[0]
    assert "--excel" not in calls[0]
    assert not (tmp_path / "out.xlsx").exists()
