"""Copias diarias de nombres y doce listas de precios; no reconstruye precios pasados."""
from __future__ import annotations
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import os
import re
import tempfile
import pandas as pd
from openpyxl import Workbook, load_workbook
from rentabilidad.infra.sql_server import SqlServerConfig, fetch_dataframe

COLOMBIA = timezone(timedelta(hours=-5))
VIEW = '[SiigoCat].[dbo].[vw_productos_activos]'
PRICE_COLUMNS = ['NOMBRE'] + [f'LISTA_PRECIO{i}' for i in range(1, 13)]
PRICE_QUERY = 'SELECT ' + ', '.join(f'[{name}]' for name in PRICE_COLUMNS) + f' FROM {VIEW}'
PATTERN = re.compile(r'^productos-(\d{4}-\d{2}-\d{2})\.xlsx$')


def snapshot_path(folder: Path, day: date) -> Path:
    return Path(folder) / f'productos-{day.isoformat()}.xlsx'


def month_cutoff(today: date) -> date:
    year, month = (today.year - 1, 12) if today.month == 1 else (today.year, today.month - 1)
    return date(year, month, min(today.day, monthrange(year, month)[1]))


def prune_snapshots(folder: Path, today: date) -> list[Path]:
    removed = []
    cutoff = month_cutoff(today)
    for path in Path(folder).glob('productos-*.xlsx'):
        match = PATTERN.fullmatch(path.name)
        if not match: continue
        try: day = date.fromisoformat(match.group(1))
        except ValueError: continue
        if day < cutoff:
            try: read_snapshot(path, day)
            except Exception: continue
            path.unlink()
            removed.append(path)
    return removed


def project_prices(data: pd.DataFrame) -> pd.DataFrame:
    names = {str(col).strip().upper(): col for col in data.columns}
    desc = next((names[name] for name in ('NOMBRE','DESCRIPCION','DESCRIPCIÓN','DESCRIPCIONINV') if name in names), None)
    prices = []
    for i in range(1, 13):
        col = next((names[name] for name in (f'LISTA_PRECIO{i}', f'PRECIO{i}') if name in names), None)
        if col is None:
            raise RuntimeError(f'Falta la lista de precio {i}; no se sustituyen columnas por posicion.')
        prices.append(col)
    if desc is None:
        raise RuntimeError('Falta la columna NOMBRE del producto.')
    result = data[[desc] + prices].copy()
    result.columns = PRICE_COLUMNS
    return result


def write_price_book(data: pd.DataFrame, metadata: dict, path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'PRODUCTOS'
    for values in data.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else value for value in values])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
    info = book.create_sheet('CAPTURA')
    metadata = dict(metadata, Formato='precios-v2')
    for row in metadata.items(): info.append(row)
    info.sheet_state = 'hidden'
    try: book.save(path)
    finally: book.close()
    from rentabilidad.infra.excel_compat import normalize_excel_application
    normalize_excel_application(path)


def read_snapshot(path: Path, day: date) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'CAPTURA' not in wb.sheetnames or 'PRODUCTOS' not in wb.sheetnames:
            raise RuntimeError('El archivo no es una captura SQL diaria valida.')
        metadata = dict(wb['CAPTURA'].iter_rows(max_col=2, values_only=True))
        if metadata.get('Fecha') != day.isoformat() or metadata.get('Vista') != VIEW:
            raise RuntimeError('La captura no corresponde a la fecha o vista requerida.')
        rows = wb['PRODUCTOS'].iter_rows(values_only=True)
        if metadata.get('Formato') == 'precios-v2':
            headers = PRICE_COLUMNS
        else:
            headers = next(rows)
        data = list(rows)
        if len(data) != metadata.get('Filas') or not data:
            raise RuntimeError('La captura esta vacia o incompleta.')
        return project_prices(pd.DataFrame(data, columns=headers))
    finally:
        wb.close()


def capture_snapshot(config: SqlServerConfig, folder: Path, *, now: datetime | None = None) -> Path:
    now = now or datetime.now(COLOMBIA)
    day = now.astimezone(COLOMBIA).date()
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    destination = snapshot_path(folder, day)
    # Primera copia del dia inmutable: nunca se sustituye por precios posteriores.
    if destination.exists():
        read_snapshot(destination, day)
        prune_snapshots(folder, day)
        return destination
    if (now.astimezone(COLOMBIA).hour, now.astimezone(COLOMBIA).minute) < (20, 45):
        raise RuntimeError('La copia de cierre se captura desde las 20:45 de Colombia, despues de la carga nocturna.')
    data = fetch_dataframe(config, PRICE_QUERY)
    if data.empty or data.columns.duplicated().any():
        raise RuntimeError('La vista de productos esta vacia o sus columnas no son unicas.')
    data = project_prices(data)
    metadata = {'Fecha': day.isoformat(), 'CapturadoUTC': now.astimezone(timezone.utc).isoformat(),
                'Vista': VIEW, 'Filas': len(data)}
    fd, name = tempfile.mkstemp(prefix='.productos-', suffix='.xlsx', dir=folder)
    os.close(fd)
    temporary = Path(name)
    try:
        write_price_book(data, metadata, temporary)
        read_snapshot(temporary, day)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    prune_snapshots(folder, day)
    return destination


def read_legacy_prices(path: Path) -> pd.DataFrame:
    """Lee el listado manual: descripcion y doce listas, sin atribuir fecha real."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'CAPTURA' in wb.sheetnames and 'PRODUCTOS' in wb.sheetnames:
            metadata = dict(wb['CAPTURA'].iter_rows(max_col=2, values_only=True))
            capture_day = date.fromisoformat(metadata['Fecha'])
            return read_snapshot(path, capture_day)
        raw = list(wb.active.iter_rows(values_only=True))
        if raw and 'NOMBRE' in [str(value).strip().upper() for value in raw[0]] and any(str(value).strip().upper() == 'LISTA_PRECIO1' for value in raw[0]):
            return project_prices(pd.DataFrame(raw[1:], columns=raw[0]))
        rows = [tuple(row[:13]) for row in raw if row and row[0] is not None]
    finally:
        wb.close()
    if not rows or not any(isinstance(value, (int, float)) for row in rows for value in row[1:]):
        raise ValueError(f'Listado de precios vacio o sin precios numericos: {path}')
    return pd.DataFrame([row + (None,) * (13-len(row)) for row in rows],
                        columns=['DESCRIPCION'] + [f'PRECIO{i}' for i in range(1,13)])


def compact_snapshot(path: Path) -> bool:
    """Proyecta una captura antigua conservando fecha, precios y procedencia originales."""
    import shutil
    from uuid import uuid4
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'CAPTURA' not in wb.sheetnames or 'PRODUCTOS' not in wb.sheetnames:
            return False
        metadata = dict(wb['CAPTURA'].iter_rows(max_col=2, values_only=True))
    finally:
        wb.close()
    day = date.fromisoformat(metadata['Fecha'])
    data = read_snapshot(path, day)
    if metadata.get('Formato') == 'precios-v2': return False
    fd, name = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(fd)
    staged = Path(name)
    fresh = path.with_name(path.stem + '.nuevo-' + uuid4().hex + '.xlsx')
    try:
        write_price_book(data, metadata, staged)
        read_snapshot(staged, day)
        shutil.copyfile(staged, fresh)
        os.replace(fresh, path)
    finally:
        staged.unlink(missing_ok=True)
        fresh.unlink(missing_ok=True)
    return True
