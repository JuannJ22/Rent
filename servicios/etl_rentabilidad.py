"""Trabajo SQL independiente para SiigoBI.EtlService; compatible con el loader instalado."""
from __future__ import annotations
import argparse
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
import logging
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from rentabilidad.core.paths import SPANISH_MONTHS


@contextmanager
def lock_report(base: Path):
    with (base / '.rentabilidad-service.lock').open('a+b') as handle:
        if handle.tell() == 0:
            handle.write(b'0'); handle.flush()
        handle.seek(0)
        if os.name == 'nt':
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        try:
            yield
        finally:
            handle.seek(0)
            if os.name == 'nt':
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def sql_config():
    from rentabilidad.infra.sql_server import SqlServerConfig
    return SqlServerConfig(
        server=os.environ.get('SQL_SERVER', ''), database='SiigoRent',
        user=os.environ.get('SQL_USER'), password=os.environ.get('SQL_PASSWORD'),
        trusted_connection=os.environ.get('SQL_TRUSTED') == '1',
        driver='ODBC Driver 18 for SQL Server', encrypt=True, trust_server_certificate=False,
    )


def check_sql(day: date):
    import pyodbc
    config = sql_config()
    if not config.server:
        raise RuntimeError('Falta SQL_SERVER para rentabilidad.')
    connection = pyodbc.connect(config.connection_string())
    try:
        cursor = connection.cursor()
        try:
            row = cursor.execute('SELECT TOP (1) 1 FROM dbo.vw_rentabilidad_cliente WHERE FECHA = ?', day.isoformat()).fetchone()
            if row is None:
                raise RuntimeError('SQL no contiene rentabilidad para la fecha solicitada; no se publica un informe vacio.')
        finally:
            cursor.close()
    finally:
        connection.close()


def validate_report(path: Path, day: date):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = wb.worksheets[0]
        expected = f'{SPANISH_MONTHS[day.month].upper()} {day.day:02d}'
        if sheet.title != expected:
            raise RuntimeError('El nombre de la hoja no corresponde a la fecha solicitada.')
        found = any(isinstance(row[5], (int, float)) and isinstance(row[6], (int, float))
                    for row in sheet.iter_rows(min_row=7, max_col=7, values_only=True))
        if not found:
            raise RuntimeError('El informe no contiene filas numericas de ventas y costos.')
        if any(len(s.title) > 31 for s in wb):
            raise RuntimeError('El informe contiene un nombre de hoja no valido para Excel.')
        if 'CCOSTO 4' in wb.sheetnames:
            zone = wb['CCOSTO 4']
            rows = zone.iter_rows(min_row=2, max_col=7, values_only=True)
            details = total = False
            for row in rows:
                if isinstance(row[0], str) and row[0].strip().upper() == 'TOTAL TIENDA PINTUCO (ZONA 7)':
                    total = True
                elif isinstance(row[3], (int, float)):
                    details = True
            if details and not total:
                raise RuntimeError('Falta el total SQL de CCOSTO 4.')
    finally:
        wb.close()


def generate(base: Path, template: Path, day: date, timeout: int, replace=False):
    base.mkdir(parents=True, exist_ok=True)
    destination = base / 'InformesDiarios' / str(day.year) / SPANISH_MONTHS[day.month] / f'{SPANISH_MONTHS[day.month]} {day.day:02d}.xlsx'
    with lock_report(base):
        if destination.exists() and not replace:
            validate_report(destination, day)
            logging.info('Informe existente validado: %s', destination)
            return destination
        if not template.is_file():
            raise FileNotFoundError(f'No existe la plantilla: {template}')
        from rentabilidad.infra.product_snapshots import snapshot_path, read_snapshot
        folder = base / 'Productos'
        snapshot = snapshot_path(folder, day)
        legacy_prices = os.environ.get("SQL_LEGACY_PRODUCT_FILE")
        if legacy_prices:
            from rentabilidad.infra.product_snapshots import read_legacy_prices
            read_legacy_prices(Path(legacy_prices))
            logging.warning("Listado manual compartido; fecha real de precios no verificada: %s", legacy_prices)
        elif not snapshot.exists():
            raise RuntimeError(f'Falta la copia de productos de {day}: {snapshot}. No se usan precios actuales para una fecha pasada.')
        if not legacy_prices:
            read_snapshot(snapshot, day)
        check_sql(day)
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.rent-', dir=destination.parent) as stage:
            staged = Path(stage) / destination.name
            shutil.copyfile(template, staged)
            command = [sys.executable, str(ROOT / 'hojas' / 'hoja01_loader.py'), '--sql',
                       '--sql-driver', 'ODBC Driver 18 for SQL Server', '--fecha', day.isoformat(), '--excel', str(staged)]
            result = subprocess.run(command, cwd=ROOT, env=dict(os.environ, PYTHONUTF8='1', PYTHONIOENCODING='utf-8', SQL_PRODUCT_SNAPSHOT=str(snapshot)),
                                    capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=timeout)
            for output in (result.stdout, result.stderr):
                if output.strip(): logging.info('%s', output.strip())
            if result.returncode:
                raise RuntimeError(f'El motor rentabilidad fallo (exit={result.returncode}).')
            if legacy_prices:
                wb = load_workbook(staged)
                name = "ORIGEN_PRECIOS"
                if name in wb.sheetnames:
                    del wb[name]
                ws = wb.create_sheet(name)
                ws.append(["Fecha informe", day.isoformat()])
                ws.append(["Listado utilizado", Path(legacy_prices).name])
                ws.append(["Advertencia", "Listado manual compartido: precios historicos del dia no verificados."])
                wb.save(staged)
                wb.close()
            validate_report(staged, day)
            os.replace(staged, destination)
    logging.info('Informe publicado: %s', destination)
    return destination


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--fecha', help='YYYY-MM-DD; predeterminado ayer en Colombia')
    parser.add_argument('--base-dir', default=r'C:\Rentabilidad')
    parser.add_argument('--template', default=r'C:\Rentabilidad\PLANTILLA.xlsx')
    parser.add_argument('--timeout', type=int, default=1800)
    parser.add_argument('--replace', action='store_true')
    parser.add_argument('--check-sql', action='store_true', help='Valida SQL y la vista de productos sin generar archivos')
    parser.add_argument('--capture-products', action='store_true', help='Guarda la vista actual de productos con la fecha real de captura')
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s', stream=sys.stdout)
    try:
        if args.timeout <= 0: raise ValueError('Timeout debe ser positivo.')
        if args.capture_products:
            from rentabilidad.infra.product_snapshots import capture_snapshot, COLOMBIA
            if args.fecha and date.fromisoformat(args.fecha) != datetime.now(COLOMBIA).date():
                raise ValueError('No se puede guardar la vista actual como si fuera una fecha pasada.')
            base = Path(args.base_dir)
            base.mkdir(parents=True, exist_ok=True)
            with lock_report(base):
                saved = capture_snapshot(sql_config(), base / 'Productos')
            logging.info('Copia diaria de productos guardada: %s', saved)
            return 0
        day = date.fromisoformat(args.fecha) if args.fecha else (datetime.now(timezone(timedelta(hours=-5))).date() - timedelta(days=1))
        os.environ['RENT_DIR'] = args.base_dir
        if args.check_sql:
            check_sql(day)
            from rentabilidad.infra.sql_server import fetch_dataframe
            if fetch_dataframe(sql_config(), 'SELECT TOP (1) * FROM [SiigoCat].[dbo].[vw_productos_activos]').empty:
                raise RuntimeError('La vista de productos esta vacia.')
            logging.info('SQL y vista de productos: correctos. No se han generado archivos.')
            return 0
        generate(Path(args.base_dir), Path(args.template), day, args.timeout, args.replace)
        return 0
    except Exception as exc:
        logging.error('Rentabilidad no completada: %s', exc)
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
